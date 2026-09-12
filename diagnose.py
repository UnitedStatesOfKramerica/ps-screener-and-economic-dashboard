"""
One-off diagnostic -- NOT part of the permanent codebase. Run once, paste
the full output back.

Inspects the real structure of SPY's daily holdings file (published free by
State Street, SPY's issuer) before writing any parsing code against it --
same reasoning as the FRED/Stooq/yfinance checks earlier: don't guess at an
unfamiliar file's column layout, look at it first.
"""
import io
import requests
import openpyxl

URL = "https://www.ssga.com/us/en/institutional/library-content/products/fund-data/etfs/us/holdings-daily-us-en-spy.xlsx"

r = requests.get(URL, timeout=30, headers={"User-Agent": "Mozilla/5.0"})
print(f"HTTP {r.status_code}, {len(r.content)} bytes, content-type={r.headers.get('content-type')}")

wb = openpyxl.load_workbook(io.BytesIO(r.content), read_only=True, data_only=True)
print("sheet names:", wb.sheetnames)

ws = wb.active
print(f"\nactive sheet: {ws.title}")
for i, row in enumerate(ws.iter_rows(values_only=True)):
    print(f"  row {i}: {row}")
    if i >= 16:
        print("  ... (truncated)")
        break
